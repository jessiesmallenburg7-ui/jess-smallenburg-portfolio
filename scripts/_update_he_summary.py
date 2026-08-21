"""Restyle HE summary DOCX, tighten section 6 to shipped prototype updates, export PDF."""
from __future__ import annotations

import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, Twips

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "projects/clinician-flow/reports/Lumen_Chart_HE_Summary_Report.docx"
PDF = ROOT / "projects/clinician-flow/reports/Lumen_Chart_HE_Summary_Report.pdf"
FONT = "Calibri"

# Absolute sizes after +2–3pt bump (keep compact enough for ≤3 pages)
SIZE_MAP = {
    18.0: 20.0,  # title
    14.0: 16.0,  # subtitle
    13.0: 15.0,  # H2
    11.0: 13.0,  # body emphasis
    10.0: 12.0,  # body / meta
    9.0: 11.0,  # footer
}


def set_run_font(run, size_pt: float | None = None, bold: bool | None = None) -> None:
    run.font.name = FONT
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), FONT)
    rFonts.set(qn("w:hAnsi"), FONT)
    rFonts.set(qn("w:cs"), FONT)
    rFonts.set(qn("w:eastAsia"), FONT)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold


def bump_size(pt: float | None, fallback: float) -> float:
    if pt is None:
        return fallback
    return SIZE_MAP.get(round(pt, 1), pt + 2.0)


def style_paragraph(p, *, fallback_size: float, bold: bool | None = None) -> None:
    # Prefer existing run size; else fallback
    existing = None
    for r in p.runs:
        if r.font.size:
            existing = r.font.size.pt
            break
    size = bump_size(existing, fallback_size)
    if not p.runs:
        run = p.add_run(p.text)
        # clear direct text on p if any (rare)
        set_run_font(run, size, bold)
        return
    for r in p.runs:
        set_run_font(r, size, bold if bold is not None else r.font.bold)


def set_spacing(p, before=60, after=40, line=1.08) -> None:
    pf = p.paragraph_format
    pf.space_before = Twips(before)
    pf.space_after = Twips(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line


def clear_runs(p) -> None:
    for r in list(p.runs):
        r._element.getparent().remove(r._element)


def set_text(p, text: str, *, size: float, bold: bool = False) -> None:
    clear_runs(p)
    run = p.add_run(text)
    set_run_font(run, size, bold)


def force_doc_defaults(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(12)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), FONT)
    rFonts.set(qn("w:hAnsi"), FONT)
    rFonts.set(qn("w:cs"), FONT)
    rFonts.set(qn("w:eastAsia"), FONT)
    for name in ("Heading 1", "Heading 2", "Heading 3", "Title", "Subtitle"):
        try:
            s = doc.styles[name]
            s.font.name = FONT
            rPr = s.element.get_or_add_rPr()
            rFonts = rPr.get_or_add_rFonts()
            rFonts.set(qn("w:ascii"), FONT)
            rFonts.set(qn("w:hAnsi"), FONT)
            rFonts.set(qn("w:cs"), FONT)
            rFonts.set(qn("w:eastAsia"), FONT)
        except KeyError:
            pass
    # Slightly tighter margins so +2–3pt type still fits ≤3 pages
    for section in doc.sections:
        section.top_margin = Twips(720)  # 0.5"
        section.bottom_margin = Twips(720)
        section.left_margin = Twips(720)
        section.right_margin = Twips(720)


def rewrite_section_6(doc: Document) -> None:
    """Keep only prototype updates that actually shipped."""
    # Find heading paragraph
    heading_i = None
    for i, p in enumerate(doc.paragraphs):
        if "Relationship to Moderated Usability Testing" in p.text:
            heading_i = i
            break
    if heading_i is None:
        raise RuntimeError("Section 6 heading not found")

    # Expected structure: heading + 4 body paras (62-65). Rewrite in place.
    bodies = [
        (
            "This heuristic evaluation was conducted separately from one moderated "
            "think-aloud session on the same prototype. Complementary findings informed "
            "a single refinement pass; only the updates that shipped in the live "
            "prototype are listed below."
        ),
        (
            "From the moderated session: a skippable first-load onboarding wizard that "
            "clarifies amber AI drafts, edit affordances, and per-section Accept with "
            "two-gate signing."
        ),
        (
            "From heuristic review (and shipped): Unaccept / re-review on SOAP sections; "
            "a strengthened pre-sign gate with locked Sign until both acknowledgments are "
            "checked and specific guidance when one is missing; a persistent visit-context "
            "sticky across SOAP and Pre-sign; plus keyboard shortcuts, clearer aria-live "
            "announcements, and remaining-section feedback on the disabled Sign control."
        ),
        (
            "Full before/after detail is documented on the Lumen Chart case study site."
        ),
    ]

    # Ensure we have enough following paragraphs; replace text of next 4 non-empty or create by overwriting known indices
    # Current doc: heading_i, then 4 body paras, then end/footer
    next_paras = []
    j = heading_i + 1
    while j < len(doc.paragraphs) and len(next_paras) < 4:
        if doc.paragraphs[j].text.strip() or len(next_paras) > 0:
            # take contiguous block after heading until we have 4 or hit End of Summary
            t = doc.paragraphs[j].text.strip()
            if t.startswith("— End") or t.startswith("- End"):
                break
            next_paras.append(doc.paragraphs[j])
        j += 1

    set_text(
        doc.paragraphs[heading_i],
        "6. Relationship to Moderated Usability Testing",
        size=15,
        bold=True,
    )
    set_spacing(doc.paragraphs[heading_i], before=120, after=60, line=1.05)

    if len(next_paras) < 4:
        raise RuntimeError(
            f"Expected 4 body paragraphs after section 6 heading, found {len(next_paras)}"
        )
    for idx, text in enumerate(bodies):
        p = next_paras[idx]
        set_text(p, text, size=12, bold=False)
        set_spacing(p, before=40, after=40, line=1.08)


def restyle_all(doc: Document) -> None:
    force_doc_defaults(doc)

    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        style_name = p.style.name if p.style else "Normal"

        if i == 0:  # title
            style_paragraph(p, fallback_size=20, bold=True)
            set_spacing(p, before=0, after=60, line=1.05)
            continue
        if i == 1:  # subtitle
            style_paragraph(p, fallback_size=16, bold=True)
            set_spacing(p, before=0, after=40, line=1.05)
            continue
        if i in (2, 3):  # meta
            style_paragraph(p, fallback_size=12)
            set_spacing(p, before=0, after=20, line=1.05)
            continue

        if style_name.startswith("Heading") or (
            text[:2].isdigit() and ". " in text[:6]
        ) or text in {
            "Method",
            "1. Severity Summary",
            "2. Top Positive Findings (Severity 0)",
            "3. Highest-Priority Findings (Severity 3 · Major)",
            "4. Priority Recommendations",
            "5. Evaluation Method (Brief)",
        }:
            # Section headings — force Calibri + bumped size
            style_paragraph(p, fallback_size=15, bold=True)
            set_spacing(p, before=120, after=50, line=1.05)
            continue

        if style_name == "List Bullet":
            style_paragraph(p, fallback_size=12)
            set_spacing(p, before=20, after=20, line=1.05)
            continue

        if text.startswith("— End") or text.startswith("Prepared for"):
            style_paragraph(p, fallback_size=11)
            set_spacing(p, before=40, after=20, line=1.05)
            continue

        # Finding labels / titles often 10–11pt
        style_paragraph(p, fallback_size=12)
        set_spacing(p, before=30, after=30, line=1.08)

    # Tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    style_paragraph(p, fallback_size=11)
                    set_spacing(p, before=20, after=20, line=1.0)

    # Explicitly fix the H5 finding title that had no font (serif fallback in PDF)
    for p in doc.paragraphs:
        if "Pre-sign acknowledgments could be clicked through" in p.text:
            style_paragraph(p, fallback_size=13, bold=True)
        if p.text.strip().startswith("H5"):
            style_paragraph(p, fallback_size=12, bold=True)


def find_soffice() -> str | None:
    for c in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ):
        if Path(c).exists():
            return c
    for c in ("soffice", "libreoffice"):
        which = shutil.which(c)
        if which:
            return which
    return None


def export_pdf() -> None:
    out_dir = DOCX.parent
    soffice = find_soffice()
    if soffice:
        cmd = [
            soffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(out_dir),
            str(DOCX),
        ]
        print("Running:", " ".join(cmd))
        subprocess.run(cmd, check=True)
        print("PDF written:", PDF)
        return

    print("No LibreOffice; trying Word COM via powershell...")
    ps = f"""
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$doc = $word.Documents.Open('{str(DOCX)}')
$doc.SaveAs([ref] '{str(PDF)}', [ref] 17)
$doc.Close()
$word.Quit()
"""
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
    print("PDF written via Word:", PDF)


def main() -> int:
    doc = Document(str(DOCX))
    restyle_all(doc)
    rewrite_section_6(doc)
    # Re-apply font on section 6 after rewrite (already set in rewrite)
    doc.save(str(DOCX))
    print("Saved DOCX:", DOCX)

    # Verify section 6 text
    doc2 = Document(str(DOCX))
    for i, p in enumerate(doc2.paragraphs):
        if "Relationship to Moderated" in p.text or (
            i > 0
            and "Relationship to Moderated" in doc2.paragraphs[max(0, i - 4) : i + 1][0].text
        ):
            pass
    in_s6 = False
    for p in doc2.paragraphs:
        if "Relationship to Moderated" in p.text:
            in_s6 = True
        if in_s6:
            print("S6:", p.text[:120])
            if p.text.startswith("— End") or p.text.startswith("Prepared"):
                break
            if "Full before/after" in p.text:
                # one more after may be end
                continue

    export_pdf()
    return 0


if __name__ == "__main__":
    sys.exit(main())
